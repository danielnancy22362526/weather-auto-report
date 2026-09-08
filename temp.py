import os
import requests
import pandas as pd
from datetime import datetime

import matplotlib
matplotlib.use("Agg")  # 背景繪圖，防止介面卡住
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.font_manager as fm

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 字型設定（避免 Linux 報錯）
plt.rcParams['font.sans-serif'] = ['DejaVu Sans', 'Arial', 'Helvetica']
plt.rcParams['axes.unicode_minus'] = False
my_font = fm.FontProperties()

# 四個觀測地點設定
LOCATIONS = [
    {"id": "Huilong", "name_tw": "新北迴龍", "name_en": "Huilong (New Taipei)", "color": "#E62117", "bar_color": "#FF9999", "lat": 25.0216, "lon": 121.4116, "tz": "Asia/Taipei"},
    {"id": "Taichung", "name_tw": "台中", "name_en": "Taichung", "color": "#FF7F0E", "bar_color": "#FFBB78", "lat": 24.1477, "lon": 120.6736, "tz": "Asia/Taipei"},
    {"id": "Chicago", "name_tw": "美國芝加哥", "name_en": "Chicago (USA)", "color": "#1F77B4", "bar_color": "#AEC7E8", "lat": 41.8781, "lon": -87.6298, "tz": "America/Chicago"},
    {"id": "Schweinfurt", "name_tw": "德國Schweinfurt", "name_en": "Schweinfurt (Germany)", "color": "#2CA02C", "bar_color": "#98DF8A", "lat": 50.0494, "lon": 10.2334, "tz": "Europe/Berlin"},
]

def get_weather_desc(rain_mm):
    if rain_mm >= 2.5:
        return f"大雨 ({rain_mm:.1f})"
    elif rain_mm > 0:
        return f"微雨 ({rain_mm:.1f})"
    else:
        return "無雨 (0.0)"

def fetch_7days_data(loc):
    url = "https://api.open-meteo.com/v1/forecast"
    params = {
        "latitude": loc["lat"],
        "longitude": loc["lon"],
        "hourly": ["temperature_2m", "relative_humidity_2m", "rain"],
        "daily": ["sunrise", "sunset"],
        "forecast_days": 7,
        "timezone": loc["tz"],
    }
    r = requests.get(url, params=params, timeout=15)
    r.raise_for_status()
    data = r.json()

    daily_sun = {}
    for d_str, sr, ss in zip(data["daily"]["time"], data["daily"]["sunrise"], data["daily"]["sunset"]):
        daily_sun[d_str] = {"sunrise": sr.split("T")[-1], "sunset": ss.split("T")[-1]}

    records = []
    for t_str, t_val, h_val, r_val in zip(data["hourly"]["time"], data["hourly"]["temperature_2m"], data["hourly"]["relative_humidity_2m"], data["hourly"]["rain"]):
        d_part, h_part = t_str.split("T")
        sun_info = daily_sun.get(d_part, {"sunrise": "-", "sunset": "-"})
        records.append({
            "地點": loc["name_tw"],
            "完整時間": t_str.replace("T", " "),
            "日期": d_part,
            "小時": h_part[:2],
            "氣溫 (°C)": t_val,
            "相對濕度 (%)": h_val,
            "降雨量 (mm)": r_val,
            "雨況說明": get_weather_desc(r_val),
            "日出時間": sun_info["sunrise"],
            "日落時間": sun_info["sunset"]
        })
    return pd.DataFrame(records)

def generate_single_chart_en(df, loc_en_name, color, bar_color, start_date, end_date, img_path):
    df["dt"] = pd.to_datetime(df["完整時間"])
    fig, ax1 = plt.subplots(figsize=(15, 6), dpi=150)
    fig.patch.set_facecolor('white')
    ax1.set_facecolor('white')

    ax2 = ax1.twinx()
    max_rain = max(df["降雨量 (mm)"].max(), 5.0)
    ax2.set_ylim(0, max_rain * 3.5)
    ax2.bar(df["dt"], df["降雨量 (mm)"], width=0.035, color=bar_color, zorder=2)
    ax2.set_ylabel("Precipitation (mm)", color="#1F77B4", fontsize=11, fontweight="bold")
    ax2.tick_params(axis='y', colors="#1F77B4")

    t_min = df["氣溫 (°C)"].min() - 3
    t_max = df["氣溫 (°C)"].max() + 5
    ax1.set_ylim(t_min, t_max)
    ax1.plot(df["dt"], df["氣溫 (°C)"], color=color, linewidth=2.2, zorder=4)
    ax1.scatter(df["dt"], df["氣溫 (°C)"], color=color, s=14, zorder=5)
    ax1.set_ylabel("Temperature (°C)", color=color, fontsize=11, fontweight="bold")
    ax1.tick_params(axis='y', colors=color)

    ax1.xaxis.set_major_locator(mdates.DayLocator())
    ax1.xaxis.set_major_formatter(mdates.DateFormatter('%m/%d\n(%a)'))
    ax1.xaxis.set_minor_locator(mdates.HourLocator(byhour=[6, 12, 18]))
    ax1.grid(True, which="major", axis="x", color="#B0B0B0", linestyle="-", linewidth=0.9)
    ax1.grid(True, which="minor", axis="x", color="#EBEBEB", linestyle=":", linewidth=0.6)
    ax1.grid(True, which="major", axis="y", color="#EBEBEB", linestyle="-", linewidth=0.8)

    for spine in ["top", "left", "right"]:
        ax1.spines[spine].set_visible(False)
        ax2.spines[spine].set_visible(False)
    ax1.spines["bottom"].set_color("#333333")

    plt.title(f"{loc_en_name} - 7-Day Hourly Weather Trend\n[ Forecast Range: {start_date} to {end_date} ]", fontsize=14, fontweight="bold", pad=15)
    plt.tight_layout()
    plt.savefig(img_path, bbox_inches="tight")
    plt.close(fig)

def generate_combined_chart_en(all_dfs, start_date, end_date, img_path):
    fig, ax1 = plt.subplots(figsize=(18, 7.5), dpi=150)
    fig.patch.set_facecolor('white')
    ax1.set_facecolor('white')

    ax2 = ax1.twinx()
    max_rain = max(max(df["降雨量 (mm)"].max() for df in all_dfs), 5.0)
    ax2.set_ylim(0, max_rain * 3.8)

    all_t_min = min(df["氣溫 (°C)"].min() for df in all_dfs) - 3
    all_t_max = max(df["氣溫 (°C)"].max() for df in all_dfs) + 6
    ax1.set_ylim(all_t_min, all_t_max)

    for idx, df in enumerate(all_dfs):
        loc_cfg = LOCATIONS[idx]
        dts = pd.to_datetime(df["完整時間"])
        offset = (idx - 1.5) * 0.008
        ax2.bar(dts + pd.to_timedelta(offset, unit='D'), df["降雨量 (mm)"], width=0.012, color=loc_cfg["bar_color"], alpha=0.65, zorder=2)
        ax1.plot(dts, df["氣溫 (°C)"], color=loc_cfg["color"], linewidth=2.2, label=f"{loc_cfg['name_en']} Temp", zorder=4)

    ax1.set_ylabel("Temperature (°C)", fontsize=12, fontweight="bold", color="#333333")
    ax2.set_ylabel("Precipitation (mm)", fontsize=12, fontweight="bold", color="#555555")

    ax1.xaxis.set_major_locator(mdates.DayLocator())
    ax1.xaxis.set_major_formatter(mdates.DateFormatter('%m/%d\n(%a)'))
    ax1.xaxis.set_minor_locator(mdates.HourLocator(byhour=[6, 12, 18]))
    ax1.grid(True, which="major", axis="x", color="#A0A0A0", linestyle="-", linewidth=0.9)
    ax1.grid(True, which="minor", axis="x", color="#EBEBEB", linestyle=":", linewidth=0.6)
    ax1.grid(True, which="major", axis="y", color="#EBEBEB", linestyle="-", linewidth=0.8)

    for spine in ["top", "left", "right"]:
        ax1.spines[spine].set_visible(False)
        ax2.spines[spine].set_visible(False)
    ax1.spines["bottom"].set_color("#333333")

    ax1.legend(loc="upper left", frameon=True, facecolor="#F9F9F9", edgecolor="#DDDDDD", fontsize=10)
    plt.title(f"4-City Combined 7-Day Weather Comparison (Hourly)\n[ Forecast Period: {start_date} to {end_date} ]", fontsize=15, fontweight="bold", pad=15)
    plt.tight_layout()
    plt.savefig(img_path, bbox_inches="tight")
    plt.close(fig)

def run_process():
    today_str = datetime.now().strftime("%Y-%m-%d")
    excel_filename = f"四城市一週氣象對比表_{today_str}.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    temp_images = []
    all_dfs = []

    border_style = Border(left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'), top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD'))
    center_align = Alignment(horizontal="center", vertical="center")

    for loc in LOCATIONS:
        name_tw = loc["name_tw"]
        name_en = loc["name_en"]
        df = fetch_7days_data(loc)
        all_dfs.append(df)

        start_date = df["日期"].iloc[0]
        end_date = df["日期"].iloc[-1]

        ws = wb.create_sheet(title=name_tw.replace(" ", "_")[:30])
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = f"【{name_tw}】一週逐時氣象明細表（日期：{start_date} 至 {end_date}）"
        title_cell.font = Font(name="DejaVu Sans", size=13, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1F4E79", fill_type="solid")
        title_cell.alignment = center_align
        ws.row_dimensions[1].height = 32

        headers = ["完整時間", "日期", "小時", "氣溫 (°C)", "相對濕度 (%)", "降雨量 (mm)", "日出時間", "日落時間"]
        ws.append(headers)
        ws.row_dimensions[2].height = 24

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=2, column=c_idx)
            c.font = Font(name="DejaVu Sans", bold=True)
            c.fill = PatternFill(start_color="D9E1F2", fill_type="solid")
            c.alignment = center_align

        for r_idx, r in df.iterrows():
            ws.append([r["完整時間"], r["日期"], f"{r['小時']}:00", r["氣溫 (°C)"], r["相對濕度 (%)"], r["降雨量 (mm)"], r["日出時間"], r["日落時間"]])
            cur_row = r_idx + 3
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=cur_row, column=c_idx)
                cell.alignment = center_align
                cell.border = border_style

        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, len(df) + 3))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        img_name = f"chart_{loc['id']}_{today_str}.png"
        generate_single_chart_en(df, name_en, loc["color"], loc["bar_color"], start_date, end_date, img_name)
        temp_images.append(img_name)

        chart_img = OpenpyxlImage(img_name)
        chart_img.width = 880
        chart_img.height = 420
        ws.add_image(chart_img, "J2")

    ws_comp = wb.create_sheet(title="一週四城市對比總表", index=0)
    start_d = all_dfs[0]["日期"].iloc[0]
    end_d = all_dfs[0]["日期"].iloc[-1]

    ws_comp.merge_cells("A1:Q1")
    comp_title = ws_comp["A1"]
    comp_title.value = f"四大地點一週（168 小時）逐時氣象綜合對比總表（日期：{start_d} 至 {end_d}）"
    comp_title.font = Font(name="DejaVu Sans", size=14, bold=True, color="FFFFFF")
    comp_title.fill = PatternFill(start_color="0D233A", fill_type="solid")
    comp_title.alignment = center_align
    ws_comp.row_dimensions[1].height = 34

    ws_comp.merge_cells("A2:A3")
    ws_comp["A2"] = "時間 (日期 HH:00)"
    ws_comp["A2"].alignment = center_align
    ws_comp["A2"].font = Font(name="DejaVu Sans", bold=True)
    ws_comp["A2"].fill = PatternFill(start_color="EAECEE", fill_type="solid")

    city_colors = ["C6EFCE", "BDD7EE", "FCE4D6", "FFF2CC"]
    col_ptr = 2
    for idx, loc in enumerate(LOCATIONS):
        start_c = col_ptr
        end_c = col_ptr + 3
        ws_comp.merge_cells(start_row=2, start_column=start_c, end_row=2, end_column=end_c)
        c_top = ws_comp.cell(row=2, column=start_c, value=loc["name_tw"])
        c_top.font = Font(name="DejaVu Sans", bold=True, size=11)
        c_top.alignment = center_align
        c_top.fill = PatternFill(start_color=city_colors[idx], fill_type="solid")

        sub_headers = ["氣溫 (°C)", "雨況說明", "日出時間", "日落時間"]
        for i, sub in enumerate(sub_headers):
            sub_c = ws_comp.cell(row=3, column=start_c + i, value=sub)
            sub_c.font = Font(name="DejaVu Sans", bold=True, size=9)
            sub_c.alignment = center_align
            sub_c.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            sub_c.border = border_style
        col_ptr += 4

    num_hours = len(all_dfs[0])
    for h_idx in range(num_hours):
        time_label = all_dfs[0].iloc[h_idx]["完整時間"]
        row_vals = [time_label]
        for df_item in all_dfs:
            r = df_item.iloc[h_idx]
            row_vals.extend([r["氣溫 (°C)"], r["雨況說明"], r["日出時間"], r["日落時間"]])
        ws_comp.append(row_vals)
        cur_row = h_idx + 4
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws_comp.cell(row=cur_row, column=c_idx)
            cell.alignment = center_align
            cell.border = border_style

    for col_idx in range(1, len(row_vals) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws_comp.cell(row=r, column=col_idx).value or "")) for r in range(2, num_hours + 4))
        ws_comp.column_dimensions[col_letter].width = max(max_len + 3, 12)

    combined_img_name = f"chart_combined_4cities_en_{today_str}.png"
    generate_combined_chart_en(all_dfs, start_d, end_d, combined_img_name)
    temp_images.append(combined_img_name)

    comb_img = OpenpyxlImage(combined_img_name)
    comb_img.width = 1050
    comb_img.height = 460
    ws_comp.add_image(comb_img, "S2")

    wb.save(excel_filename)
    print(f"\n[成功] 報表已產出：{excel_filename}", flush=True)

    for f in temp_images:
        if os.path.exists(f):
            os.remove(f)

if __name__ == "__main__":
    run_process()
